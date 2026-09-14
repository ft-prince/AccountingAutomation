"""File → grid → header row → dict rows. Real bank exports carry preamble lines (account
holder, period, "****" separators) above the table, so the header is found, not assumed."""

import csv
import io
import re
from html.parser import HTMLParser
from typing import Any

OLE2_MAGIC = b"\xd0\xcf\x11\xe0"
PDF_MAGIC = b"%PDF"
HEADER_SCAN_ROWS = 60
FORMATS = ("csv", "xlsx", "xls", "pdf")


class StatementError(ValueError):
    pass


def normalise_header(h: Any) -> str:
    """'Withdrawal Amount (INR )' and 'withdrawal amount(inr)' are the same column."""
    s = re.sub(r"\s+", " ", str(h or "")).strip().lower()
    return re.sub(r"\s*([()./])\s*", r"\1", s)


def format_for(filename: str, data: bytes) -> str:
    name = filename.lower()
    if data[:4] == PDF_MAGIC or name.endswith(".pdf"):
        return "pdf"
    if name.endswith(".xlsx"):
        return "xlsx"
    if name.endswith(".xls"):
        return "xls"
    return "csv"


# ---------------------------------------------------------------- grids


def _csv_grid(text: str) -> list[list[str]]:
    sample = "\n".join(text.splitlines()[:HEADER_SCAN_ROWS])
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [list(r) for r in csv.reader(io.StringIO(text), dialect)]


def _xlsx_grid(data: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
    return [list(r) for r in ws.iter_rows(values_only=True)]


class _TableParser(HTMLParser):
    """Banks label HTML tables as .xls; the stdlib parser is enough for <tr><td>."""

    def __init__(self) -> None:
        super().__init__()
        self.rows: list[list[str]] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: Any) -> None:
        if tag == "tr":
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []

    def handle_endtag(self, tag: str) -> None:
        if tag in ("td", "th") and self._row is not None and self._cell is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            self.rows.append(self._row)
            self._row = None

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)


def _xls_grid(data: bytes) -> list[list[Any]]:
    if data[:4] == OLE2_MAGIC:
        raise StatementError(
            "Legacy binary .xls is not supported — download the statement as CSV, XLSX or PDF."
        )
    text = data.decode("utf-8-sig", errors="replace")
    if text.lstrip()[:1] == "<":
        parser = _TableParser()
        parser.feed(text)
        return [list(r) for r in parser.rows]
    return _csv_grid(text)


LINE_TOLERANCE = 3.0  # points; words whose tops differ less are one row
CELL_GAP = 8.0  # points; a wider gap between words starts a new header cell


def _lines(page: Any) -> list[list[dict[str, Any]]]:
    words = sorted(page.extract_words(), key=lambda w: (round(w["top"]), w["x0"]))
    lines: list[list[dict[str, Any]]] = []
    for w in words:
        if lines and abs(lines[-1][0]["top"] - w["top"]) <= LINE_TOLERANCE:
            lines[-1].append(w)
        else:
            lines.append([w])
    return [sorted(line, key=lambda w: w["x0"]) for line in lines]


def _cells(line: list[dict[str, Any]]) -> list[tuple[float, float, str]]:
    """Merge words separated by a normal space into one cell: (x0, x1, text)."""
    cells: list[tuple[float, float, str]] = []
    for w in line:
        if cells and w["x0"] - cells[-1][1] < CELL_GAP:
            x0, _, text = cells[-1]
            cells[-1] = (x0, w["x1"], f"{text} {w['text']}")
        else:
            cells.append((w["x0"], w["x1"], w["text"]))
    return cells


def _words_grid(page: Any, required: list[tuple[str, str]]) -> list[list[Any]]:
    """Unruled PDF: locate the header line, then bucket every word below it by column x-range.
    ponytail: a right-aligned number wider than its header can spill into the previous column;
    move to explicit per-bank column x-ranges if a real statement needs it."""
    lines = _lines(page)
    for i, line in enumerate(lines):
        header = _cells(line)
        names = {normalise_header(c[2]) for c in header}
        if not any(d in names and s in names for d, s in required):
            continue
        bounds = [(header[k][1] + header[k + 1][0]) / 2 for k in range(len(header) - 1)]
        grid: list[list[Any]] = [[c[2] for c in header]]
        for row in lines[i + 1 :]:
            cells = [""] * len(header)
            for w in row:
                centre = (w["x0"] + w["x1"]) / 2
                col = sum(1 for b in bounds if centre > b)
                cells[col] = f"{cells[col]} {w['text']}".strip()
            grid.append(cells)
        return grid
    return []


def _pdf_grid(
    data: bytes, password: str | None, required: list[tuple[str, str]]
) -> list[list[Any]]:
    import pdfplumber
    from pdfminer.pdfdocument import PDFPasswordIncorrect
    from pdfplumber.utils.exceptions import PdfminerException

    grid: list[list[Any]] = []
    try:
        with pdfplumber.open(io.BytesIO(data), password=password or "") as pdf:
            for page in pdf.pages:
                tables = page.extract_tables() or [_words_grid(page, required)]
                for table in tables:
                    grid.extend([" ".join(str(c or "").split()) for c in row] for row in table)
    except PdfminerException as exc:
        if exc.args and isinstance(exc.args[0], PDFPasswordIncorrect):
            raise StatementError(
                "This PDF is password-protected — enter the password the bank set "
                "(usually customer ID or date of birth)."
            ) from exc
        raise StatementError(f"Could not read PDF: {exc}") from exc
    if not grid:
        raise StatementError(
            "No table found in this PDF. Scanned statements are not supported — "
            "download CSV/XLSX from net banking instead."
        )
    return grid


def read_grid(
    data: bytes,
    fmt: str,
    *,
    password: str | None = None,
    required: list[tuple[str, str]] | None = None,
) -> list[list[Any]]:
    """`required` lists (date, description) normalised header pairs used to spot the header."""
    if fmt == "csv":
        return _csv_grid(data.decode("utf-8-sig", errors="replace"))
    if fmt == "xlsx":
        return _xlsx_grid(data)
    if fmt == "xls":
        return _xls_grid(data)
    if fmt == "pdf":
        return _pdf_grid(data, password, required or [])
    raise StatementError(f"unsupported format {fmt}")


# ---------------------------------------------------------------- header detection


def find_header(grid: list[list[Any]], required: list[tuple[str, str]]) -> int | None:
    """Index of the first row (within the scan window) holding both columns of any mapping."""
    for i, row in enumerate(grid[:HEADER_SCAN_ROWS]):
        cells = {normalise_header(c) for c in row}
        if any(d in cells and s in cells for d, s in required):
            return i
    return None


def grid_to_rows(grid: list[list[Any]], header_index: int) -> list[dict[str, Any]]:
    header = [normalise_header(h) for h in grid[header_index]]
    out: list[dict[str, Any]] = []
    for row in grid[header_index + 1 :]:
        if not any(c not in (None, "") for c in row):
            continue
        out.append(dict(zip(header, row, strict=False)))
    return out
