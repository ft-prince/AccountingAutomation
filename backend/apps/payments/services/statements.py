"""Bank statement import: per-bank column mappings, sha256 per row, signed amounts."""

import csv
import hashlib
import io
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from django.db import IntegrityError, transaction

from apps.payments.models import (
    BankAccount,
    BankBalanceSnapshot,
    BankStatementImport,
    BankTransaction,
)


@dataclass(frozen=True)
class ColumnMapping:
    key: str
    date: str
    description: str
    debit: str | None
    credit: str | None
    amount: str | None  # single signed column
    reference: str | None
    balance: str | None
    date_formats: tuple[str, ...]


# Column headers as they appear in the bank's own export. Registry, not if-statements.
MAPPINGS: dict[str, ColumnMapping] = {
    "hdfc": ColumnMapping(
        "hdfc",
        "Date",
        "Narration",
        "Withdrawal Amt.",
        "Deposit Amt.",
        None,
        "Chq./Ref.No.",
        "Closing Balance",
        ("%d/%m/%y", "%d/%m/%Y"),
    ),
    "icici": ColumnMapping(
        "icici",
        "Transaction Date",
        "Transaction Remarks",
        "Withdrawal Amount (INR )",
        "Deposit Amount (INR )",
        None,
        "Cheque Number",
        "Balance (INR )",
        ("%d/%m/%Y",),
    ),
    "sbi": ColumnMapping(
        "sbi",
        "Txn Date",
        "Description",
        "Debit",
        "Credit",
        None,
        "Ref No./Cheque No.",
        "Balance",
        ("%d %b %Y", "%d-%m-%Y"),
    ),
    "generic": ColumnMapping(
        "generic",
        "date",
        "description",
        "debit",
        "credit",
        "amount",
        "reference",
        "balance",
        ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"),
    ),
}


class StatementError(ValueError):
    pass


def _dec(v: Any) -> Decimal | None:
    if v is None:
        return None
    s = str(v).strip().replace(",", "").replace("₹", "")
    if not s or s in {"-", "–"}:
        return None
    try:
        return Decimal(s)
    except InvalidOperation as exc:
        raise StatementError(f"not a number: {v!r}") from exc


def _date(v: Any, formats: tuple[str, ...]) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for f in formats:
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            continue
    raise StatementError(f"unparseable date: {v!r}")


def read_rows(data: bytes, fmt: str) -> list[dict[str, Any]]:
    if fmt == "csv":
        text = data.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        return [{(k or "").strip(): v for k, v in r.items()} for r in reader]
    if fmt == "xlsx":
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(h).strip() if h is not None else "" for h in rows[0]]
        return [
            dict(zip(header, r, strict=False)) for r in rows[1:] if any(c is not None for c in r)
        ]
    raise StatementError(f"unsupported format {fmt}")


def detect_mapping(headers: set[str]) -> ColumnMapping:
    for m in MAPPINGS.values():
        if m.date in headers and m.description in headers:
            return m
    raise StatementError(f"No bank mapping matches columns: {sorted(headers)}")


@dataclass(frozen=True)
class ParsedRow:
    date: date
    amount: Decimal
    description: str
    reference: str
    balance_after: Decimal | None
    sha256: str


def parse_rows(
    rows: list[dict[str, Any]], mapping: ColumnMapping, account_id: str
) -> list[ParsedRow]:
    out: list[ParsedRow] = []
    for r in rows:
        d = _date(r.get(mapping.date), mapping.date_formats)
        desc = str(r.get(mapping.description) or "").strip()
        ref = str(r.get(mapping.reference) or "").strip() if mapping.reference else ""
        if mapping.amount and r.get(mapping.amount) not in (None, ""):
            amount = _dec(r.get(mapping.amount)) or Decimal("0")
        else:
            debit = _dec(r.get(mapping.debit)) if mapping.debit else None
            credit = _dec(r.get(mapping.credit)) if mapping.credit else None
            amount = (credit or Decimal("0")) - (debit or Decimal("0"))
        bal = _dec(r.get(mapping.balance)) if mapping.balance else None
        key = f"{account_id}|{d.isoformat()}|{amount}|{desc}|{ref}"
        out.append(
            ParsedRow(
                d, amount, desc[:500], ref[:100], bal, hashlib.sha256(key.encode()).hexdigest()
            )
        )
    return out


def import_statement(
    account: BankAccount,
    *,
    data: bytes,
    filename: str,
    fmt: str,
    actor: Any = None,
    mapping_key: str | None = None,
) -> BankStatementImport:
    rows = read_rows(data, fmt)
    if not rows:
        raise StatementError("No rows found.")
    mapping = MAPPINGS[mapping_key] if mapping_key else detect_mapping(set(rows[0].keys()))
    parsed = parse_rows(rows, mapping, str(account.pk))
    with transaction.atomic():
        imp = BankStatementImport.objects.create(
            bank_account=account,
            filename=filename[:255],
            format=fmt,
            mapping=mapping.key,
            rows_total=len(parsed),
            created_by=actor if getattr(actor, "pk", None) else None,
        )
        imported = dup = 0
        latest: ParsedRow | None = None
        for p in parsed:
            try:
                with transaction.atomic():
                    BankTransaction.objects.create(
                        bank_account=account,
                        statement_import=imp,
                        date=p.date,
                        amount=p.amount,
                        description=p.description,
                        reference=p.reference,
                        balance_after=p.balance_after,
                        sha256=p.sha256,
                    )
                imported += 1
            except IntegrityError:
                dup += 1
            if p.balance_after is not None and (latest is None or p.date >= latest.date):
                latest = p
        imp.rows_imported, imp.rows_duplicate = imported, dup
        imp.save(update_fields=["rows_imported", "rows_duplicate", "updated_at"])
        if latest is not None:
            BankBalanceSnapshot.objects.get_or_create(
                org=account.org,
                bank_account=account,
                date=latest.date,
                source="statement",
                defaults={"balance": latest.balance_after},
            )
    return imp
