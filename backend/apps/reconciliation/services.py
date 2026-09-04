"""Orchestration, transactions, side effects: import a GSTR-2B download, run matching,
override a match, record an IMS action (PROJECT_SPECS §3.8, §7.2)."""

import json
import re
from calendar import monthrange
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from typing import Any

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from apps.core.audit import record as audit
from apps.invoices.models import Direction, Invoice, InvoiceStatus
from apps.reconciliation.domain.matching import (
    MATCH_TYPES,
    MatchSide,
    itc_at_risk,
    match_records,
)
from apps.reconciliation.domain.parsers import (
    Parse2BError,
    Parsed2BRecord,
    parse_2b_json,
    parse_2b_rows,
)
from apps.reconciliation.models import (
    BatchSource,
    GSTR2BBatch,
    GSTR2BRecord,
    MatchType,
    ReconciliationMatch,
)

PERIOD_PATTERN = re.compile(r"^(0[1-9]|1[0-2])(\d{4})$")
B2B_SHEET = "B2B"
# Books invoices dated up to this many days outside the 2B period are still candidates.
PERIOD_SLACK = timedelta(days=3)
ZERO = Decimal("0")


class ReconciliationError(ValueError):
    """User-facing failure (bad upload, bad period, conflicting override)."""


def period_bounds(period: str) -> tuple[date, date]:
    match = PERIOD_PATTERN.match(period or "")
    if match is None:
        raise ReconciliationError(f"period must be MMYYYY, got {period!r}")
    month, year = int(match.group(1)), int(match.group(2))
    return date(year, month, 1), date(year, month, monthrange(year, month)[1])


def detect_source(filename: str, data: bytes) -> str:
    lower = filename.lower()
    if lower.endswith(".json"):
        return BatchSource.JSON
    if lower.endswith(".xlsx"):
        return BatchSource.XLSX
    if data[:2] == b"PK":
        return BatchSource.XLSX
    if data.lstrip()[:1] == b"{":
        return BatchSource.JSON
    raise ReconciliationError("upload must be a GSTN GSTR-2B .json or .xlsx file")


def _sheet_rows(data: bytes) -> list[list[object]]:
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a zoo of types for a bad zip
        raise ReconciliationError("could not open the workbook") from exc
    sheet = next((workbook[n] for n in workbook.sheetnames if n.upper() == B2B_SHEET), None)
    if sheet is None:
        raise ReconciliationError(f"workbook has no {B2B_SHEET!r} sheet")
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def parse_upload(data: bytes, source: str) -> tuple[str, list[Parsed2BRecord]]:
    """(period declared in the file or "", records). Raises ReconciliationError."""
    try:
        if source == BatchSource.JSON:
            try:
                payload = json.loads(data)
            except ValueError as exc:
                raise ReconciliationError("upload is not valid JSON") from exc
            return parse_2b_json(payload)
        return "", parse_2b_rows(_sheet_rows(data))
    except Parse2BError as exc:
        raise ReconciliationError(str(exc)) from exc


@transaction.atomic
def import_2b(
    org: Any, *, data: bytes, filename: str, period: str, actor: Any = None
) -> GSTR2BBatch:
    period_bounds(period)
    source = detect_source(filename, data)
    declared, parsed = parse_upload(data, source)
    if declared and declared != period:
        raise ReconciliationError(f"file is for period {declared}, not {period}")
    if not parsed:
        raise ReconciliationError("no B2B invoices found in the upload")
    batch = GSTR2BBatch.objects.create(
        org=org,
        period=period,
        source=source,
        filename=filename[:255],
        imported_by=actor if actor is not None and actor.is_authenticated else None,
        imported_at=timezone.now(),
        counts={"records": len(parsed)},
    )
    GSTR2BRecord.objects.bulk_create(
        GSTR2BRecord(
            org=org,
            batch=batch,
            supplier_gstin=p.supplier_gstin,
            supplier_name=p.supplier_name[:200],
            invoice_number=p.invoice_number[:32],
            invoice_date=p.invoice_date,
            invoice_value=p.invoice_value,
            place_of_supply=p.place_of_supply,
            reverse_charge=p.reverse_charge,
            taxable_value=p.taxable_value,
            igst=p.igst,
            cgst=p.cgst,
            sgst=p.sgst,
            cess=p.cess,
            itc_available=p.itc_available,
            raw=p.raw,
        )
        for p in parsed
    )
    audit(
        org,
        actor=actor,
        entity=batch,
        action="reconciliation.import_2b",
        after={"period": period, "source": source, "records": len(parsed)},
    )
    return batch


def record_tax(rec: GSTR2BRecord) -> Decimal:
    return rec.igst + rec.cgst + rec.sgst + rec.cess


def invoice_tax(inv: Invoice) -> Decimal:
    return inv.igst + inv.cgst + inv.sgst + inv.cess


def _record_side(rec: GSTR2BRecord) -> MatchSide:
    return MatchSide(
        str(rec.pk),
        rec.supplier_gstin,
        rec.invoice_number,
        rec.invoice_date,
        rec.invoice_value,
        record_tax(rec),
    )


def _invoice_side(inv: Invoice) -> MatchSide:
    return MatchSide(
        str(inv.pk),
        (inv.party.gstin or "").upper(),
        inv.invoice_number,
        inv.invoice_date,
        inv.total,
        invoice_tax(inv),
    )


def candidate_invoices(batch: GSTR2BBatch) -> list[Invoice]:
    start, end = period_bounds(batch.period)
    return list(
        Invoice.objects.for_org(batch.org)
        .filter(
            direction=Direction.INWARD,
            status=InvoiceStatus.CONFIRMED,
            invoice_date__range=(start - PERIOD_SLACK, end + PERIOD_SLACK),
            party__gstin__isnull=False,
        )
        .exclude(party__gstin="")
        .select_related("party")
        .order_by("invoice_date", "invoice_number")
    )


def match_at_risk(match: ReconciliationMatch) -> Decimal:
    tax = record_tax(match.record) if match.record_id and match.record else ZERO
    return itc_at_risk(match.match_type, tax, match.delta_tax)


def compute_counts(batch: GSTR2BBatch) -> dict[str, Any]:
    matches = list(batch.matches.select_related("record"))
    counts: dict[str, Any] = {t: 0 for t in MATCH_TYPES}
    total_risk = ZERO
    for m in matches:
        counts[m.match_type] = counts.get(m.match_type, 0) + 1
        total_risk += match_at_risk(m)
    counts["records"] = batch.records.count()
    counts["itc_at_risk"] = str(total_risk)
    return counts


@transaction.atomic
def run_reconciliation(batch: GSTR2BBatch, *, actor: Any = None) -> GSTR2BBatch:
    """Recompute every match for the batch. Matches are derived data: a re-run replaces
    them (including manual overrides — the override endpoint exists for after the run)."""
    start, end = period_bounds(batch.period)
    invoices = candidate_invoices(batch)
    records = list(batch.records.all())
    by_invoice = {str(i.pk): i for i in invoices}
    by_record = {str(r.pk): r for r in records}
    period_keys = frozenset(k for k, i in by_invoice.items() if start <= i.invoice_date <= end)
    result = match_records(
        tuple(_record_side(r) for r in records),
        tuple(_invoice_side(i) for i in invoices),
        period_keys,
    )
    batch.matches.all().delete()
    rows = [
        ReconciliationMatch(
            org=batch.org,
            batch=batch,
            record=by_record[p.record_key],
            invoice=by_invoice[p.invoice_key],
            match_type=p.outcome.match_type,
            delta_value=p.outcome.delta_value,
            delta_tax=p.outcome.delta_tax,
        )
        for p in result.pairs
    ]
    rows += [
        ReconciliationMatch(
            org=batch.org, batch=batch, record=by_record[k], match_type=MatchType.MISSING_IN_BOOKS
        )
        for k in result.missing_in_books
    ]
    rows += [
        ReconciliationMatch(
            org=batch.org, batch=batch, invoice=by_invoice[k], match_type=MatchType.MISSING_IN_2B
        )
        for k in result.missing_in_2b
    ]
    ReconciliationMatch.objects.bulk_create(rows)
    batch.counts = compute_counts(batch)
    batch.save(update_fields=["counts", "updated_at"])
    audit(batch.org, actor=actor, entity=batch, action="reconciliation.run", after=batch.counts)
    return batch


def _match_snapshot(match: ReconciliationMatch) -> dict[str, Any]:
    return {
        "match_type": match.match_type,
        "invoice": str(match.invoice_id) if match.invoice_id else None,
        "note": match.note,
        "delta_value": str(match.delta_value),
        "delta_tax": str(match.delta_tax),
    }


def _absorb_placeholder(match: ReconciliationMatch, invoice: Invoice) -> None:
    """Linking an invoice supersedes its missing_in_2b placeholder; a row that already
    pairs it with another 2B record is a conflict the reviewer must undo first."""
    holder = (
        ReconciliationMatch.objects.filter(batch=match.batch, invoice=invoice)
        .exclude(pk=match.pk)
        .first()
    )
    if holder is None:
        return
    if holder.record_id:
        raise ReconciliationError("that invoice is already matched in this batch")
    holder.delete()


def _release_invoice(batch: GSTR2BBatch, invoice: Invoice) -> None:
    """An invoice unlinked by an override is back to missing_in_2b when it is in-period."""
    start, end = period_bounds(batch.period)
    if start <= invoice.invoice_date <= end:
        ReconciliationMatch.objects.create(
            org=batch.org, batch=batch, invoice=invoice, match_type=MatchType.MISSING_IN_2B
        )


@transaction.atomic
def override_match(
    match: ReconciliationMatch,
    *,
    actor: Any,
    match_type: str | None = None,
    note: str | None = None,
    invoice: Invoice | None = None,
    clear_invoice: bool = False,
) -> ReconciliationMatch:
    before = _match_snapshot(match)
    released: Invoice | None = None
    if invoice is not None:
        _absorb_placeholder(match, invoice)
        released = match.invoice if match.invoice_id != invoice.pk else None
        match.invoice = invoice
        if match_type is None and match.match_type == MatchType.MISSING_IN_BOOKS:
            match_type = MatchType.FUZZY  # a manual link is a fuzzy match unless told otherwise
    elif clear_invoice:
        if not match.record_id:
            raise ReconciliationError("a missing_in_2b row has no record to keep; re-run instead")
        released = match.invoice
        match.invoice = None
        if match_type is None:
            match_type = MatchType.MISSING_IN_BOOKS
    if match_type is not None:
        match.match_type = match_type
    if note is not None:
        match.note = note[:500]
    if match.record_id and match.invoice_id and match.record and match.invoice:
        match.delta_value = match.record.invoice_value - match.invoice.total
        match.delta_tax = record_tax(match.record) - invoice_tax(match.invoice)
    else:
        match.delta_value = ZERO
        match.delta_tax = ZERO
    match.resolved_by = actor if actor is not None and actor.is_authenticated else None
    match.resolved_at = timezone.now()
    match.save()
    if released is not None:
        _release_invoice(match.batch, released)
    batch = match.batch
    batch.counts = compute_counts(batch)
    batch.save(update_fields=["counts", "updated_at"])
    audit(
        match.org,
        actor=actor,
        entity=match,
        action="reconciliation.override",
        before=before,
        after=_match_snapshot(match),
    )
    return match


@transaction.atomic
def set_ims_action(rec: GSTR2BRecord, *, action: str, note: str, actor: Any) -> GSTR2BRecord:
    """IMS accept / reject / pend on one 2B record (§3.8). Append-only audit trail."""
    before = {"ims_action": rec.ims_action, "ims_note": rec.ims_note}
    rec.ims_action = action
    rec.ims_note = note[:500]
    rec.save(update_fields=["ims_action", "ims_note", "updated_at"])
    audit(
        rec.org,
        actor=actor,
        entity=rec,
        action=f"ims.{action}",
        before=before,
        after={"ims_action": rec.ims_action, "ims_note": rec.ims_note},
    )
    return rec
